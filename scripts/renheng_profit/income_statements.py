"""利润表文件选择、解析、标准化及 M3 抵销前汇总。"""
from __future__ import annotations

import csv
from dataclasses import asdict
from decimal import Decimal
import json
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook

from .common import (
    COMPANIES,
    COMPANY_FILE_KEYWORDS,
    ITEM_ALIASES,
    KEY_ITEMS,
    PRE_SUMMARY_FILENAME,
    SUMMARY_CUMULATIVE_COLUMN,
    SUMMARY_MONTHLY_COLUMN,
    US_COMPANY,
    ZERO,
    ParseMetadata,
    SourceFile,
    compact_text,
    ignore_us_profit_for_period,
    json_default,
    parse_profit_period,
    target_month_from_period_token,
    to_number,
)


def register_sources(
    source_dir: Path,
    report_year: str,
    period_token: str | None,
    optional_companies: frozenset[str] = frozenset(),
    excluded_companies: frozenset[str] = frozenset(),
) -> list[SourceFile]:
    """从包含多类输入资料的目录中，识别目标年度六家主体的利润表文件。"""
    registered: list[SourceFile] = []
    errors: list[str] = []
    report_label = period_token or f"{report_year}年"
    target_month = target_month_from_period_token(period_token)

    for company in COMPANIES:
        if company in excluded_companies:
            continue
        if company == US_COMPANY and ignore_us_profit_for_period(report_year, target_month):
            continue
        candidates = [
            path for path in source_dir.rglob("*")
            if path.is_file()
            and path.suffix.lower() in {".xls", ".xlsx"}
            and "报表" in compact_text(path.name)
            and report_year in compact_text(path.name)
            and any(keyword in compact_text(path.name) for keyword in COMPANY_FILE_KEYWORDS[company])
        ]
        if company in {US_COMPANY, "新疆公司"}:
            eligible = [
                (parse_profit_period(path.name, report_year), path)
                for path in candidates
            ]
            eligible = [
                (period_month, path)
                for period_month, path in eligible
                if period_month is not None
                and target_month is not None
                and period_month <= target_month
            ]
            if eligible:
                latest_month = max(period_month for period_month, _ in eligible)
                matches = [path for period_month, path in eligible if period_month == latest_month]
            else:
                matches = []
        else:
            matches = [
                path for path in candidates
                if parse_profit_period(path.name, report_year) == target_month
            ]
        if len(matches) != 1:
            if not matches and company in optional_companies:
                continue
            detail = "未找到" if not matches else f"找到{len(matches)}份"
            errors.append(
                f"{company}的{report_label}利润表{detail}；同一目标期间只能有一份有效文件"
            )
            continue
        path = matches[0]
        registered.append(SourceFile(company, str(path), path.suffix.lower(), path.stat().st_size))
    if errors:
        raise FileNotFoundError("；".join(errors))
    return registered


def read_sheet_rows(path: Path) -> Iterable[tuple[str, list[list[Any]]]]:
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return
    workbook = xlrd.open_workbook(path)
    for sheet_name in workbook.sheet_names():
        sheet = workbook.sheet_by_name(sheet_name)
        yield sheet_name, [
            [sheet.cell_value(row_index, col_index) for col_index in range(sheet.ncols)]
            for row_index in range(sheet.nrows)
        ]


def find_profit_sheet(path: Path) -> tuple[str, list[list[Any]]]:
    """以工作表名或前五行标题为“利润表”作为稳定识别规则。"""
    candidates: list[tuple[str, list[list[Any]]]] = []
    for sheet_name, rows in read_sheet_rows(path):
        has_profit_title = compact_text(sheet_name) == "利润表" or any(
            compact_text(cell) == "利润表" for row in rows[:5] for cell in row
        )
        if has_profit_title:
            candidates.append((sheet_name, rows))
    if len(candidates) != 1:
        detail = "未找到" if not candidates else f"找到{len(candidates)}张"
        raise ValueError(f"{path.name}{detail}标题为“利润表”的工作表")
    return candidates[0]


def find_header_columns(
    rows: list[list[Any]], file_name: str, sheet_name: str
) -> tuple[int, int, int, int]:
    """利润表取数前的必需列校验：项目、本月金额、本年累计金额缺一不可。"""
    required_headers = {
        "项目列": {"项目", "项目名称"},
        "本月金额列": {"本月金额", "本月数", "本期金额", "本季度金额", "本季度数", "本季金额"},
        "本年累计金额列": {"本年累计金额", "本年金额", "本年累计数"},
    }
    closest_missing: list[str] = list(required_headers)

    for row_index, row in enumerate(rows[:20]):
        headers = [compact_text(value) for value in row]
        found = {
            label: next((i for i, header in enumerate(headers) if header in aliases), None)
            for label, aliases in required_headers.items()
        }
        missing = [label for label, column in found.items() if column is None]
        if len(missing) < len(closest_missing):
            closest_missing = missing
        if not missing:
            columns = list(found.values())
            if len(set(columns)) != 3:
                raise ValueError(
                    f"{file_name}的“{sheet_name}”中项目、本月、累计列位置重复，停止取数"
                )
            return row_index, found["项目列"], found["本月金额列"], found["本年累计金额列"]

    missing_text = "、".join(closest_missing)
    raise ValueError(
        f"{file_name}的“{sheet_name}”缺少必需列：{missing_text}；"
        "必须同时存在项目、本月金额、本年累计金额列，停止取数"
    )


def canonical_item(raw_item: str) -> str | None:
    value = compact_text(raw_item)
    for canonical, aliases in ITEM_ALIASES.items():
        if any(alias in value for alias in aliases):
            return canonical
    return None


def parse_income_statement(source: SourceFile) -> tuple[ParseMetadata, dict[str, dict[str, Any]]]:
    path = Path(source.path)
    sheet_name, rows = find_profit_sheet(path)
    header_row, project_col, monthly_col, cumulative_col = find_header_columns(
        rows, path.name, sheet_name
    )
    normalized: dict[str, dict[str, Any]] = {}
    current_parent_item: str | None = None
    current_parent_indent: int | None = None
    for row_index, row in enumerate(rows[header_row + 1:], start=header_row + 1):
        if project_col >= len(row):
            continue
        raw_label = "" if row[project_col] is None else str(row[project_col])
        raw_item = compact_text(raw_label)
        leading_spaces = len(raw_label) - len(raw_label.lstrip())
        item = canonical_item(raw_item)
        if not item:
            continue
        # 主项目会重置层级；未识别的“其中”明细行不重置上级。
        if item != "研发费用":
            current_parent_item = item
            current_parent_indent = leading_spaces if item == "管理费用" else None
        monthly = to_number(row[monthly_col]) if monthly_col < len(row) else None
        cumulative = to_number(row[cumulative_col]) if cumulative_col < len(row) else None
        if item in KEY_ITEMS and (monthly is None or cumulative is None):
            raise ValueError(
                f"{source.company}的{path.name}第{row_index + 1}行关键科目“{item}”金额为空或无法解析"
            )
        if item == "研发费用":
            is_management_detail = (
                "研究费用" in raw_item
                and current_parent_item == "管理费用"
                and current_parent_indent is not None
                and leading_spaces > current_parent_indent
            )
            if "研发费用" in raw_item:
                source_type = "独立研发费用列"
            elif is_management_detail:
                source_type = "管理费用明细-研究费用"
            else:
                raise ValueError(
                    f"{source.company}第{row_index + 1}行研究费用未标识为管理费用明细，无法确认费用归属"
                )
            existing = normalized.get(item)
            # 独立研发费用行优先于管理费用明细行，避免同一主体重复计入。
            if existing is None or source_type == "独立研发费用列":
                normalized[item] = {
                    "monthly": ZERO if monthly is None else monthly,
                    "cumulative": ZERO if cumulative is None else cumulative,
                    "source_label": raw_item,
                    "source_type": source_type,
                    "source_row": row_index + 1,
                }
            continue
        # 投资收益主项和“其中”明细分别取数并累计，避免任一行覆盖另一行。
        if item == "投资收益":
            current_monthly = ZERO if monthly is None else monthly
            current_cumulative = ZERO if cumulative is None else cumulative
            existing = normalized.get(item)
            if existing is None:
                normalized[item] = {
                    "monthly": current_monthly,
                    "cumulative": current_cumulative,
                    "source_label": raw_item,
                    "source_row": row_index + 1,
                    "source_rows": [row_index + 1],
                }
            else:
                existing["monthly"] += current_monthly
                existing["cumulative"] += current_cumulative
                existing["source_label"] = f"{existing['source_label']} + {raw_item}"
                existing.setdefault("source_rows", [existing["source_row"]]).append(row_index + 1)
            continue
        if item not in normalized or monthly is not None or cumulative is not None:
            normalized[item] = {
                "monthly": ZERO if monthly is None else monthly,
                "cumulative": ZERO if cumulative is None else cumulative,
                "source_label": raw_item,
                "source_row": row_index + 1,
            }

    missing_key_items = [item for item in KEY_ITEMS if item not in normalized]
    if missing_key_items:
        raise ValueError(
            f"{source.company}的{path.name}缺少关键科目：{'、'.join(missing_key_items)}"
        )
    management = normalized.get("管理费用")
    research = normalized.get("研发费用")
    if management and research and research.get("source_type") == "管理费用明细-研究费用":
        raw_monthly = management["monthly"]
        raw_cumulative = management["cumulative"]
        research_monthly = research["monthly"]
        research_cumulative = research["cumulative"]

        management["original_monthly"] = raw_monthly
        management["original_cumulative"] = raw_cumulative
        if source.company == US_COMPANY and raw_monthly == ZERO and raw_cumulative == ZERO:
            # 美国利润表主项目为0而“其中”明细有数时，明细不作为独立研发费用取数。
            research["original_monthly"] = research_monthly
            research["original_cumulative"] = research_cumulative
            research["monthly"] = ZERO
            research["cumulative"] = ZERO
            research["source_type"] = "美国管理费用为0-研究费用明细不取数"
            research["ignored_reason"] = "美国管理费用本月和累计均为0"
            management["r_and_d_deducted"] = False
        else:
            management["monthly"] = raw_monthly - research_monthly
            management["cumulative"] = raw_cumulative - research_cumulative
            management["r_and_d_deducted"] = True
    metadata = ParseMetadata(
        company=source.company, file=source.path, sheet=sheet_name,
        header_row=header_row + 1, project_column=project_col + 1,
        monthly_column=monthly_col + 1, cumulative_column=cumulative_col + 1,
        source_row_count=len(rows), recognized_item_count=len(normalized),
    )
    return metadata, normalized


def aggregate_six_companies(
    parsed: dict[str, dict[str, dict[str, Any]]],
    quarterly_cumulative_companies: frozenset[str] = frozenset(),
    invoice_revenue_overrides: dict[str, dict[str, Decimal]] | None = None,
    cumulative_only_companies: frozenset[str] = frozenset(),
) -> list[dict[str, Any]]:
    """按同一科目汇总；季末将美国季度金额视为该季末月份金额；非季末仅累计纳入的主体本月置0。"""
    rows: list[dict[str, Any]] = []
    for item in ITEM_ALIASES:
        result: dict[str, Any] = {"项目": item}
        month_total = ZERO
        cumulative_total = ZERO
        for company in COMPANIES:
            values = parsed.get(company, {}).get(item, {"monthly": ZERO, "cumulative": ZERO})
            if item == "营业收入" and invoice_revenue_overrides and company in invoice_revenue_overrides:
                monthly = invoice_revenue_overrides[company]["monthly"]
                cumulative = invoice_revenue_overrides[company]["cumulative"]
            else:
                monthly = values["monthly"]
                cumulative = values["cumulative"]
            if company in cumulative_only_companies:
                # 非季末：该主体本月不计入汇总，累计保留（累计口径始终纳入美国/新疆）。
                monthly = ZERO
            result[f"{company}_本月"] = monthly
            result[f"{company}_累计"] = cumulative
            month_total += monthly
            cumulative_total += cumulative
        result[SUMMARY_MONTHLY_COLUMN] = month_total
        result[SUMMARY_CUMULATIVE_COLUMN] = cumulative_total
        rows.append(result)

    # M3只做同科目相加。源表利润行如存在差异，应提示财务复核，不在合并阶段自动修正。
    # 营业成本仍直接汇总源表数据；M5仅抵销营业收入。
    return rows


def write_profit_outputs(output_dir: Path, report_label: str, sources: list[SourceFile], metadata: list[ParseMetadata], parsed: dict[str, dict[str, dict[str, Any]]], summary: list[dict[str, Any]]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "01_输入文件登记.json").write_text(json.dumps([asdict(x) for x in sources], ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "02_利润表解析与标准化.json").write_text(json.dumps({"parse_metadata": [asdict(x) for x in metadata], "normalized_data": parsed}, ensure_ascii=False, indent=2, default=json_default), encoding="utf-8")
    with (output_dir / PRE_SUMMARY_FILENAME).open("w", newline="", encoding="utf-8-sig") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=list(summary[0]))
        writer.writeheader()
        writer.writerows(summary)
    report = [f"# 六主体利润表汇总测试结果（{report_label}）", "", "范围：步骤1输入登记、步骤2利润表解析、步骤3标准化、步骤4六主体抵销前汇总。", "本次为只读测试，未执行内部交易抵销，也未修改任何源文件。", "", "## 已识别文件"]
    report.extend(f"- {x.company}：{Path(x.path).name}" for x in sources)
    report.extend(["", "## 关键项目汇总（元）", "", "| 项目 | 本月六主体汇总 | 本年累计六主体汇总 |", "| --- | ---: | ---: |"])
    for row in summary:
        if row["项目"] in KEY_ITEMS:
            report.append(f"| {row['项目']} | {row['六主体汇总_本月']:,.2f} | {row['六主体汇总_累计']:,.2f} |")
    report.extend(["", "## 当前口径", "- 本月、累计列按表头识别，不使用固定列号。", "- 研发费用为独立利润表行时，管理费用不调整；研究费用为管理费用明细行时，从管理费用中拆出并单列研发费用。", "- 本阶段营业成本直接汇总6家源表数据；M5仅抵销营业收入，营业成本保持不变。", "- 本文件的汇总为内部交易抵销前数据。", "- 内部交易抵销将在下一阶段按开票信息表的票级发票基础信息计算。"])
    (output_dir / "04_测试结果说明.md").write_text("\n".join(report) + "\n", encoding="utf-8")
