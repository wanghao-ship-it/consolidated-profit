#!/usr/bin/env python3
"""仁恒合并利润表的 LangGraph 编排层。

金额计算复用 ``renheng_profit`` 中的确定性函数。本模块只保存可序列化状态和
文件引用，负责节点、分支、Checkpoint、interrupt 和发布门禁；工作流由网页服务
通过 ``build_workflow()`` 调用。
"""
from __future__ import annotations

import copy
from contextvars import ContextVar
from dataclasses import asdict
from datetime import datetime
from decimal import Decimal
import hashlib
import inspect
import json
from operator import add
from pathlib import Path
import re
from typing import Any, Literal
from typing_extensions import Annotated, TypedDict

from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field

import renheng_profit as core


_ARTIFACT_REFERENCE_RESOLVER: ContextVar[Any | None] = ContextVar(
    "profit_artifact_reference_resolver", default=None
)


def _artifact_path(state: ProfitWorkflowState, key: str) -> Path:
    reference = state.get("artifacts", {}).get(key)
    if not reference:
        raise FileNotFoundError(f"工作流产物索引不存在：{key}")
    resolver = _ARTIFACT_REFERENCE_RESOLVER.get()
    if resolver is not None and str(reference).startswith("minio://"):
        return resolver(str(reference))
    return Path(reference)


def _reference_path(reference: str) -> Path:
    resolver = _ARTIFACT_REFERENCE_RESOLVER.get()
    if resolver is not None and reference.startswith("minio://"):
        return resolver(reference)
    return Path(reference)


def _wrap_node(node, artifact_persister=None, artifact_resolver=None):
    if artifact_persister is None and artifact_resolver is None:
        return node

    def wrapped(state, config):
        run_id = str(config.get("configurable", {}).get("run_id", ""))
        resolver = None
        if artifact_resolver is not None:
            resolver = lambda reference: artifact_resolver(run_id, reference)
        token = _ARTIFACT_REFERENCE_RESOLVER.set(resolver)
        try:
            parameters = inspect.signature(node).parameters
            updates = node(state, config) if len(parameters) > 1 else node(state)
            if artifact_persister is not None:
                updates = artifact_persister(run_id, updates or {})
            return updates
        finally:
            _ARTIFACT_REFERENCE_RESOLVER.reset(token)

    wrapped.__name__ = getattr(node, "__name__", "workflow_node")
    return wrapped


def merge_artifacts(left: dict[str, str], right: dict[str, str]) -> dict[str, str]:
    return {**(left or {}), **(right or {})}


class ProfitWorkflowInput(BaseModel):
    year: int
    month: int = Field(ge=1, le=12)
    source_dir: Path
    output_root: Path


class ProfitWorkflowState(TypedDict, total=False):
    year: int
    month: int
    source_dir: str
    output_root: str

    workflow_status: Literal["RUNNING", "WAITING_REVIEW", "SUCCEEDED", "FAILED"]
    review_round: int

    current_period_policy: dict[str, Any]
    prior_period_policy: dict[str, Any]

    selected_files: dict[str, str | list[str]]
    current_input_fingerprint: str
    prior_input_fingerprint: str

    run_output_dir: str
    artifacts: Annotated[dict[str, str], merge_artifacts]

    m7_differences: list[dict[str, str]]
    has_difference: bool

    subject_differences: list[dict[str, str]]
    unbilled_evidence: list[dict[str, Any]]
    difference_review_path: str | None

    corrected_revenue: dict[str, Any] | None
    review_history: Annotated[list[dict[str, Any]], add]

    formal_report: str | None
    error: dict[str, Any] | None


def determine_period_policy_values(year: int, month: int) -> dict[str, Any]:
    """把年度、月份转换为本次运行唯一的美国/新疆期间政策。"""
    historical_us_zero = year == 2025 and month <= 5
    confirmation_month = 0 if month <= 2 else ((month - 3) // 3 + 1) * 3
    us_required = not historical_us_zero and confirmation_month > 0
    include_us_monthly = us_required and month in core.QUARTER_END_MONTHS
    return {
        "year": year,
        "month": month,
        "monthly_range": [month],
        "cumulative_range": list(range(1, month + 1)),
        "us_required": us_required,
        "us_confirmation_month": confirmation_month if us_required else None,
        "us_report_period": f"Q{confirmation_month // 3}" if us_required else None,
        "include_us_monthly": include_us_monthly,
        "include_us_cumulative": us_required,
        "us_quarter_start": ((month - 1) // 3) * 3 + 1 if include_us_monthly else None,
        "ignore_us": historical_us_zero or confirmation_month == 0,
        "xinjiang_mode": "ZERO",
    }


def determine_period_policy(state: ProfitWorkflowState) -> dict[str, Any]:
    return {
        "workflow_status": "RUNNING",
        "review_round": 0,
        "current_period_policy": determine_period_policy_values(state["year"], state["month"]),
        "prior_period_policy": determine_period_policy_values(state["year"] - 1, state["month"]),
        "artifacts": {},
        "review_history": [],
        "formal_report": None,
        "error": None,
    }


def _find_template_file(source_dir: Path) -> Path:
    all_xls = [
        path for path in source_dir.rglob("*.xls")
        if path.is_file() and not path.name.startswith("~$")
    ]
    named = [path for path in all_xls if "模板" in core.compact_text(path.name)]
    candidates = named or [
        path for path in all_xls
        if not (
            "报表" in core.compact_text(path.name)
            and any(
                keyword in core.compact_text(path.name)
                for keywords in core.COMPANY_FILE_KEYWORDS.values()
                for keyword in keywords
            )
        )
    ]
    if len(candidates) != 1:
        detail = "未找到" if not candidates else f"找到{len(candidates)}份"
        raise FileNotFoundError(f"合并利润表.xls模板{detail}；来源目录中必须能唯一识别模板")
    core.validate_template(candidates[0])
    return candidates[0]


def _select_us_source(source_dir: Path, year: int, confirmation_month: int) -> core.SourceFile:
    candidates: list[Path] = []
    for path in source_dir.rglob("*"):
        name = core.compact_text(path.name)
        if (
            path.is_file()
            and path.suffix.lower() in {".xls", ".xlsx"}
            and not path.name.startswith("~$")
            and "报表" in name
            and str(year) in name
            and any(keyword in name for keyword in core.COMPANY_FILE_KEYWORDS[core.US_COMPANY])
            and core.parse_profit_period(path.name, str(year)) == confirmation_month
        ):
            candidates.append(path)
    if len(candidates) != 1:
        detail = "未找到" if not candidates else f"找到{len(candidates)}份"
        raise FileNotFoundError(
            f"{year}年美国分公司确认期间{confirmation_month}月利润表{detail}；必须唯一"
        )
    path = candidates[0]
    return core.SourceFile(core.US_COMPANY, str(path), path.suffix.lower(), path.stat().st_size)


def _select_profit_sources(
    source_dir: Path, year: int, month: int, policy: dict[str, Any]
) -> list[core.SourceFile]:
    ordinary = core.register_sources(
        source_dir,
        str(year),
        f"{year}.{month}",
        optional_companies=frozenset(),
        excluded_companies=frozenset({"新疆公司", core.US_COMPANY}),
    )
    if policy["us_required"]:
        ordinary.append(
            _select_us_source(source_dir, year, int(policy["us_confirmation_month"]))
        )
    return ordinary


def _fingerprint(paths: list[Path], policy: dict[str, Any]) -> str:
    digest = hashlib.sha256()
    digest.update(core.input_fingerprint(paths).encode("ascii"))
    digest.update(core.RULE_VERSION.encode("utf-8"))
    digest.update(json.dumps(policy, ensure_ascii=False, sort_keys=True).encode("utf-8"))
    return digest.hexdigest()


def _serialize_sources(sources: list[core.SourceFile]) -> list[str]:
    """将选中的利润表主体转换为 Checkpoint 可保存的文件路径列表。"""
    return [source.path for source in sources]


def _company_for_profit_path(path: Path) -> str:
    name = core.compact_text(path.name)
    matches = [
        company for company, keywords in core.COMPANY_FILE_KEYWORDS.items()
        if any(keyword in name for keyword in keywords)
    ]
    if len(matches) != 1:
        raise ValueError(f"无法从已选利润表文件名唯一识别主体：{path.name}")
    return matches[0]


def _sources_from_state(state: ProfitWorkflowState, key: str) -> list[core.SourceFile]:
    values = state["selected_files"][key]
    if not isinstance(values, list):
        raise TypeError(f"selected_files.{key}必须是路径列表")
    result = []
    for value in values:
        path = Path(value)
        result.append(core.SourceFile(
            _company_for_profit_path(path), str(path), path.suffix.lower(), path.stat().st_size
        ))
    return result


def _current_fingerprint_paths(state: ProfitWorkflowState) -> list[Path]:
    selected = state["selected_files"]
    return [
        *(Path(path) for path in selected["current_profit_statements"]),
        *(Path(path) for path in selected["current_invoice_files"]),
        Path(str(selected["current_sales_ledger"])),
        Path(str(selected["template_file"])),
    ]


def _prior_fingerprint_paths(state: ProfitWorkflowState) -> list[Path]:
    selected = state["selected_files"]
    return [
        *(Path(path) for path in selected["prior_profit_statements"]),
        Path(str(selected["prior_sales_ledger"])),
    ]


def preflight(state: ProfitWorkflowState) -> dict[str, Any]:
    source_dir = Path(state["source_dir"])
    output_root = Path(state["output_root"])
    if not source_dir.is_dir():
        raise FileNotFoundError(f"统一来源目录不存在或不是目录：{source_dir}")
    if output_root.exists() and not output_root.is_dir():
        raise ValueError(f"输出根路径不是目录：{output_root}")
    try:
        output_root.resolve().relative_to(source_dir.resolve())
    except ValueError:
        pass
    else:
        raise ValueError("output_root不能位于source_dir内部，避免输出被下一次预检识别为输入")

    year = state["year"]
    month = state["month"]
    prior_year = year - 1
    current_sources = _select_profit_sources(
        source_dir, year, month, state["current_period_policy"]
    )
    prior_sources = _select_profit_sources(
        source_dir, prior_year, month, state["prior_period_policy"]
    )
    core.validate_profit_sources(current_sources)
    core.validate_profit_sources(prior_sources)

    invoice_paths: list[Path] = []
    for source_company in core.REQUIRED_INVOICE_SOURCES:
        for period_month in range(1, month + 1):
            path = core.find_invoice_file(source_dir, source_company, year, period_month)
            assert path is not None
            invoice_paths.append(path)
    core.calculate_invoice_revenue_overrides(source_dir, year, month)

    current_ledger = core.read_ledger_dataset(source_dir, year, month)
    prior_ledger = core.read_ledger_dataset(source_dir, prior_year, month)
    if not current_ledger.has_invoice_number:
        raise ValueError(f"{year}年销售台账缺少数电票号码/数电发票号码列")
    if not prior_ledger.has_invoice_number:
        raise ValueError(f"{prior_year}年销售台账缺少数电票号码/数电发票号码列")
    template = _find_template_file(source_dir)

    selected: dict[str, str | list[str]] = {
        "current_profit_statements": _serialize_sources(current_sources),
        "prior_profit_statements": _serialize_sources(prior_sources),
        "current_invoice_files": [str(path) for path in invoice_paths],
        "current_sales_ledger": str(current_ledger.file),
        "prior_sales_ledger": str(prior_ledger.file),
        "template_file": str(template),
    }
    temporary_state: ProfitWorkflowState = {**state, "selected_files": selected}
    return {
        "selected_files": selected,
        "current_input_fingerprint": _fingerprint(
            _current_fingerprint_paths(temporary_state), state["current_period_policy"]
        ),
        "prior_input_fingerprint": _fingerprint(
            _prior_fingerprint_paths(temporary_state), state["prior_period_policy"]
        ),
    }


def _dirs(state: ProfitWorkflowState) -> dict[str, Path]:
    root = Path(state["run_output_dir"])
    return {
        "root": root,
        "m1": root / "M1_资料及期间确认",
        "m2": root / "M2_单体数据及研发拆分",
        "m3": root / "M3_抵销前六主体汇总",
        "m4": root / "M4_内部销售",
        "m5": root / "M5_抵销后合并结果",
        "m6": root / "M6_指标及模板回填",
        "m7": root / "M7_销售台账校验",
    }


def _status_payload(state: ProfitWorkflowState, updates: dict[str, Any] | None = None) -> dict[str, Any]:
    values = {**state, **(updates or {})}
    artifacts = merge_artifacts(state.get("artifacts", {}), (updates or {}).get("artifacts", {}))
    history = list(state.get("review_history", []))
    if updates and "review_history" in updates:
        history.extend(updates["review_history"])
    return {
        "workflow_status": values.get("workflow_status", "RUNNING"),
        "year": values["year"],
        "month": values["month"],
        "rule_version": core.RULE_VERSION,
        "review_round": values.get("review_round", 0),
        "current_input_fingerprint": values.get("current_input_fingerprint"),
        "prior_input_fingerprint": values.get("prior_input_fingerprint"),
        "m7_differences": values.get("m7_differences", []),
        "has_difference": values.get("has_difference", False),
        "subject_differences": values.get("subject_differences", []),
        "unbilled_evidence": values.get("unbilled_evidence", []),
        "corrected_revenue": values.get("corrected_revenue"),
        "review_history": history,
        "formal_report": values.get("formal_report"),
        "error": values.get("error"),
        "artifacts": artifacts,
    }


def _write_status(state: ProfitWorkflowState, updates: dict[str, Any] | None = None) -> Path | None:
    run_dir = (updates or {}).get("run_output_dir") or state.get("run_output_dir")
    if not run_dir:
        return None
    return core.write_run_status(Path(run_dir), _status_payload(state, updates))


def prepare_run(state: ProfitWorkflowState, config) -> dict[str, Any]:
    thread_id = str(config.get("configurable", {}).get("thread_id", "")).strip()
    if not thread_id:
        raise ValueError("LangGraph运行必须提供稳定且唯一的thread_id")
    safe_thread_id = re.sub(r"[^0-9A-Za-z_.-]+", "_", thread_id)
    output_root = Path(state["output_root"])
    run_dir = output_root / f"consolidated_profit_{state['year']}_{state['month']:02d}" / safe_thread_id
    dirs = {
        name: path for name, path in _dirs({**state, "run_output_dir": str(run_dir)}).items()
        if name != "root"
    }
    for path in dirs.values():
        path.mkdir(parents=True, exist_ok=True)

    manifest_path = dirs["m1"] / "01_输入文件与期间政策.json"
    manifest_path.write_text(
        json.dumps(
            {
                "thread_id": thread_id,
                "年度": state["year"],
                "月份": state["month"],
                "规则版本": core.RULE_VERSION,
                "已选文件": {
                    key: [Path(value).name for value in values]
                    if isinstance(values, list)
                    else Path(values).name
                    for key, values in state["selected_files"].items()
                },
                "本年期间政策": state["current_period_policy"],
                "去年期间政策": state["prior_period_policy"],
                "本年输入指纹": state["current_input_fingerprint"],
                "去年输入指纹": state["prior_input_fingerprint"],
                "本年营业收入": "科技、阳原开票扣设备，汇入其他主体利润表收入，再扣内部销售",
                "去年营业收入": "境内销方台账基数-科技到美国+美国利润表收入",
                "去年最终合并报表参与计算": False,
                "新疆处理": "不读取、不校验，按0.00处理",
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    updates = {
        "run_output_dir": str(run_dir),
        "workflow_status": "RUNNING",
        "artifacts": {"m1_input_manifest": str(manifest_path)},
    }
    _write_status(state, updates)
    return updates


def _write_parse_bundle(path: Path, sources: list[core.SourceFile], source_root: Path) -> None:
    metadata: list[core.ParseMetadata] = []
    parsed: dict[str, dict[str, dict[str, Any]]] = {}
    for source in sources:
        parse_info, normalized = core.parse_income_statement(source)
        parse_info.file = Path(parse_info.file).resolve().relative_to(source_root.resolve()).as_posix()
        metadata.append(parse_info)
        parsed[source.company] = normalized
    source_rows = []
    for source in sources:
        row = asdict(source)
        row["path"] = Path(source.path).resolve().relative_to(source_root.resolve()).as_posix()
        source_rows.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "sources": source_rows,
                "metadata": [asdict(item) for item in metadata],
                "parsed": parsed,
            },
            ensure_ascii=False,
            indent=2,
            default=core.json_default,
        ),
        encoding="utf-8",
    )


_MONEY_KEYS = {
    "monthly", "cumulative", "original_monthly", "original_cumulative",
    "monthly_invoice_total", "monthly_equipment_total",
    "cumulative_invoice_total", "cumulative_equipment_total",
    "本月", "累计",
}


def _restore_money(value: Any, key: str | None = None) -> Any:
    if isinstance(value, dict):
        return {item_key: _restore_money(item, item_key) for item_key, item in value.items()}
    if isinstance(value, list):
        return [_restore_money(item) for item in value]
    if key in _MONEY_KEYS and isinstance(value, (str, int, float)):
        return Decimal(str(value))
    return value


def _load_parse_bundle(path: Path) -> tuple[list[core.SourceFile], list[core.ParseMetadata], dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return (
        [core.SourceFile(**item) for item in payload["sources"]],
        [core.ParseMetadata(**item) for item in payload["metadata"]],
        _restore_money(payload["parsed"]),
    )


def current_m2_parse(state: ProfitWorkflowState) -> dict[str, Any]:
    path = _dirs(state)["m2"] / "本年" / "01_单体数据与研发拆分.json"
    _write_parse_bundle(
        path,
        _sources_from_state(state, "current_profit_statements"),
        Path(state["source_dir"]),
    )
    return {"artifacts": {"current_m2": str(path)}}


def current_invoice_revenue(state: ProfitWorkflowState) -> dict[str, Any]:
    values = core.calculate_invoice_revenue_overrides(
        Path(state["source_dir"]), state["year"], state["month"]
    )
    path = _dirs(state)["m2"] / "本年" / "02_科技阳原开票口径.json"
    path.write_text(
        json.dumps(values, ensure_ascii=False, indent=2, default=core.json_default),
        encoding="utf-8",
    )
    return {"artifacts": {"current_invoice_revenue": str(path)}}


def _load_invoice_overrides(path: Path) -> dict[str, dict[str, Decimal]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {
        company: {key: Decimal(str(value)) for key, value in values.items()}
        for company, values in payload.items()
    }


def _calculation_settings(policy: dict[str, Any]) -> dict[str, Any]:
    include_monthly = bool(policy["include_us_monthly"])
    include_cumulative = bool(policy["include_us_cumulative"])
    return {
        "quarterly_companies": frozenset({core.US_COMPANY}) if include_monthly else frozenset(),
        "cumulative_only": frozenset({core.US_COMPANY}) if include_cumulative and not include_monthly else frozenset(),
        "monthly_group_names": core.active_group_company_names(include_monthly),
        "cumulative_group_names": core.active_group_company_names(include_cumulative),
        "quarter_start": policy["us_quarter_start"],
        "us_buyer_names": frozenset({core.GROUP_COMPANIES[core.US_COMPANY]}) if include_cumulative else frozenset(),
        "us_confirmation_month": policy["us_confirmation_month"],
    }


def current_m3_pre_elimination(state: ProfitWorkflowState) -> dict[str, Any]:
    sources, metadata, parsed = _load_parse_bundle(_artifact_path(state, "current_m2"))
    overrides = _load_invoice_overrides(_artifact_path(state, "current_invoice_revenue"))
    settings = _calculation_settings(state["current_period_policy"])
    summary = core.aggregate_six_companies(
        parsed, settings["quarterly_companies"], overrides, settings["cumulative_only"]
    )
    output_dir = _dirs(state)["m3"] / "本年"
    core.write_profit_outputs(
        output_dir, f"{state['year']}.{state['month']}", sources, metadata, parsed, summary
    )
    core.write_invoice_revenue_override_output(output_dir, overrides, parsed)
    return {"artifacts": {"current_m3": str(output_dir / core.PRE_SUMMARY_FILENAME)}}


def current_m4_internal_sales(state: ProfitWorkflowState) -> dict[str, Any]:
    settings = _calculation_settings(state["current_period_policy"])
    summary, detail = core.run_m4(
        Path(state["source_dir"]), state["year"], state["month"], _dirs(state)["m4"] / "本年",
        settings["monthly_group_names"], settings["cumulative_group_names"],
        settings["quarter_start"], settings["us_buyer_names"], settings["us_confirmation_month"],
    )
    return {"artifacts": {"current_m4_summary": str(summary), "current_m4_detail": str(detail)}}


def current_m5_final(state: ProfitWorkflowState) -> dict[str, Any]:
    path = core.run_m5(
        _artifact_path(state, "current_m3"),
        _artifact_path(state, "current_m4_summary"),
        _artifact_path(state, "current_m4_detail"),
        _dirs(state)["m5"] / "本年",
    )
    return {"artifacts": {"current_m5": str(path)}}


def prior_m2_parse(state: ProfitWorkflowState) -> dict[str, Any]:
    path = _dirs(state)["m2"] / "去年同期" / "01_单体数据与研发拆分.json"
    _write_parse_bundle(
        path,
        _sources_from_state(state, "prior_profit_statements"),
        Path(state["source_dir"]),
    )
    return {"artifacts": {"prior_m2": str(path)}}


def prior_m3_pre_summary(state: ProfitWorkflowState) -> dict[str, Any]:
    sources, metadata, parsed = _load_parse_bundle(_artifact_path(state, "prior_m2"))
    settings = _calculation_settings(state["prior_period_policy"])
    summary = core.aggregate_six_companies(
        parsed, settings["quarterly_companies"], None, settings["cumulative_only"]
    )
    output_dir = _dirs(state)["m3"] / "去年同期"
    core.write_profit_outputs(
        output_dir, f"{state['year'] - 1}.{state['month']}", sources, metadata, parsed, summary
    )
    return {"artifacts": {"prior_m3": str(output_dir / core.PRE_SUMMARY_FILENAME)}}


def prior_m4_ledger_components(state: ProfitWorkflowState) -> dict[str, Any]:
    prior_year = state["year"] - 1
    policy = state["prior_period_policy"]
    dataset = core.read_ledger_dataset(Path(state["source_dir"]), prior_year, state["month"])
    domestic_sellers = frozenset(core.GROUP_COMPANY_NAMES - {core.GROUP_COMPANIES[core.US_COMPANY]})
    domestic = core.ledger_summary(dataset, state["month"], domestic_sellers)
    confirmation = int(policy["us_confirmation_month"] or 0)
    us_internal = core.read_prior_ledger_us_internal(
        dataset.file, state["month"], confirmation, policy["us_quarter_start"]
    )
    pre_rows = core.read_pre_elimination_summary(_artifact_path(state, "prior_m3"))
    revenue_row = next(row for row in pre_rows if row["项目"] == "营业收入")
    us_revenue = {
        "本月": core.elimination_to_money(revenue_row.get(f"{core.US_COMPANY}_本月", "0")),
        "累计": core.elimination_to_money(revenue_row.get(f"{core.US_COMPANY}_累计", "0")),
    }
    payload = {
        "domestic_ledger": {"本月": domestic["本月"], "累计": domestic["累计"]},
        "us_internal": us_internal,
        "us_revenue": us_revenue,
    }
    path = _dirs(state)["m4"] / "去年同期" / "01_去年收入组成.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=core.json_default),
        encoding="utf-8",
    )
    return {"artifacts": {"prior_m4_components": str(path)}}


def prior_m5_final(state: ProfitWorkflowState) -> dict[str, Any]:
    components = _restore_money(
        json.loads(_artifact_path(state, "prior_m4_components").read_text(encoding="utf-8"))
    )
    pre_rows = core.read_pre_elimination_summary(_artifact_path(state, "prior_m3"))
    monthly = components["domestic_ledger"]["本月"] - components["us_internal"]["本月"] + components["us_revenue"]["本月"]
    cumulative = components["domestic_ledger"]["累计"] - components["us_internal"]["累计"] + components["us_revenue"]["累计"]
    rows = core.build_prior_year_rows(pre_rows, monthly, cumulative)
    core.recalculate_operating_cost_after_elimination(rows)
    path = _dirs(state)["m5"] / "去年同期" / "01_去年同期重新计算结果.csv"
    core.write_result_rows(path, rows)
    return {"artifacts": {"prior_m5": str(path)}}


def _self_sales_and_us(state: ProfitWorkflowState) -> tuple[dict[str, Any], dict[str, Any], dict[str, Decimal], dict[str, Decimal]]:
    current_settings = _calculation_settings(state["current_period_policy"])
    current_yangyuan = core.calculate_yangyuan_self_sales(
        Path(state["source_dir"]), state["year"], state["month"],
        current_settings["monthly_group_names"], current_settings["cumulative_group_names"],
    )
    prior_dataset = core.read_ledger_dataset(Path(state["source_dir"]), state["year"] - 1, state["month"])
    prior_yangyuan = core.read_prior_sales_ledger(prior_dataset.file, state["month"])
    current_us = core.read_us_xinjiang_revenue(_artifact_path(state, "current_m3"))
    prior_us = core.read_us_xinjiang_revenue(_artifact_path(state, "prior_m3"))
    return current_yangyuan, prior_yangyuan, current_us, prior_us


def m6_preview(state: ProfitWorkflowState) -> dict[str, Any]:
    current_rows = core.read_result_rows(_artifact_path(state, "current_m5"))
    prior_rows = core.read_result_rows(_artifact_path(state, "prior_m5"))
    output = core.write_m6_outputs(
        _dirs(state)["m6"], core.build_m6_output_rows(prior_rows, current_rows)
    )
    current_yangyuan, prior_yangyuan, current_us, prior_us = _self_sales_and_us(state)
    self_sales = core.write_yangyuan_self_sales_output(
        _dirs(state)["m6"], current_yangyuan, prior_yangyuan, current_us, prior_us
    )
    return {"artifacts": {"m6_preview": str(output), "self_sales": str(self_sales)}}


def _m7_inputs(state: ProfitWorkflowState) -> tuple[Any, ...]:
    source_dir = Path(state["source_dir"])
    year = state["year"]
    month = state["month"]
    current_dataset = core.read_ledger_dataset(source_dir, year, month)
    prior_dataset = core.read_ledger_dataset(source_dir, year - 1, month)
    current_ledger = core.ledger_summary(current_dataset, month)
    domestic_sellers = frozenset(core.GROUP_COMPANY_NAMES - {core.GROUP_COMPANIES[core.US_COMPANY]})
    prior_domestic = core.ledger_summary(prior_dataset, month, domestic_sellers)
    current_settings = _calculation_settings(state["current_period_policy"])
    prior_policy = state["prior_period_policy"]
    us_buyer_internal = core.read_us_buyer_internal(
        _artifact_path(state, "current_m4_detail"),
        int(current_settings["us_confirmation_month"] or 0), month,
    )
    prior_us_internal = core.read_prior_ledger_us_internal(
        prior_dataset.file, month, int(prior_policy["us_confirmation_month"] or 0),
        prior_policy["us_quarter_start"],
    )
    current_us = core.read_us_xinjiang_revenue(_artifact_path(state, "current_m3"))
    prior_us = core.read_us_xinjiang_revenue(_artifact_path(state, "prior_m3"))
    return (
        current_dataset, prior_dataset, current_ledger, prior_domestic,
        current_us, us_buyer_internal, prior_us, prior_us_internal,
    )


def m7_group_validation(state: ProfitWorkflowState) -> dict[str, Any]:
    _, _, current_ledger, prior_domestic, current_us, us_buyer_internal, prior_us, prior_us_internal = _m7_inputs(state)
    rows = core.build_m7_validation_rows(
        core.read_result_rows(_artifact_path(state, "current_m5")),
        core.read_result_rows(_artifact_path(state, "prior_m5")),
        current_ledger, prior_domestic, current_us, us_buyer_internal, prior_us, prior_us_internal,
    )
    output = core.write_m7_validation_outputs(_dirs(state)["m7"], rows)
    has_difference = bool(core.nonzero_m7_keys(rows))
    return {
        "m7_differences": rows,
        "has_difference": has_difference,
        "difference_review_path": None if not has_difference else state.get("difference_review_path"),
        "artifacts": {"m7_validation": str(output)},
    }


def route_after_m7(state: ProfitWorkflowState) -> Literal["prepare_difference_review", "publish_report"]:
    return "prepare_difference_review" if state["has_difference"] else "publish_report"


def _json_evidence(evidence: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    return [
        {**item, "金额": format(item["金额"], ".2f"), "差异ID": diff_id}
        for diff_id, values in evidence.items()
        for item in values
    ]


def write_difference_review_workbook(
    output_path: Path,
    rows: list[dict[str, str]],
    evidence_by_diff: dict[str, list[dict[str, Any]]],
    subject_rows: list[dict[str, str]],
    state: ProfitWorkflowState,
    review_round: int,
) -> None:
    workbook = Workbook()
    info = workbook.active
    info.title = "审核信息"
    info.append(["字段", "值"])
    for key, value in (
        ("审核轮次", review_round),
        ("年度", state["year"]),
        ("月份", state["month"]),
        ("本年输入指纹", state["current_input_fingerprint"]),
        ("去年输入指纹", state["prior_input_fingerprint"]),
        ("规则版本", core.RULE_VERSION),
        ("差异快照哈希", core.diff_snapshot_hash(rows)),
        ("处理说明", "审核差异证据后，通过同一thread_id提交有差异范围的正确营业收入"),
    ):
        info.append([key, value])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 80

    detail = workbook.create_sheet("差异明细")
    headers = [
        "差异ID", "年度类型", "范围", "M5可比营业收入", "台账可比收入", "差异",
        "无票销售金额", "无票后剩余差异", "无票归因", "证据源文件", "证据工作表", "证据源行号",
    ]
    detail.append(headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for row in rows:
        evidence = evidence_by_diff.get(row["差异ID"], [])
        detail.append([
            row["差异ID"], row["年度类型"], row["范围"],
            float(Decimal(row["M5抵销后营业收入"])),
            float(Decimal(row["销售台账金额（不含税）"])),
            float(Decimal(row["差异"])), float(Decimal(row["无票销售金额"])),
            float(Decimal(row["无票后剩余差异"])), row["无票归因"],
            "；".join(sorted({item["源文件"] for item in evidence})),
            "；".join(sorted({item["源工作表"] for item in evidence})),
            "；".join(str(item["源行号"]) for item in evidence),
        ])

    subject = workbook.create_sheet("主体差异拆分")
    subject.append(["范围", "主体", "主体差异", "对集团差异的影响"])
    for cell in subject[1]:
        cell.font = Font(bold=True)
    signed_number_format = "+#,##0.00;-#,##0.00;0.00"
    for item in subject_rows:
        subject.append([
            item["范围"], item["主体"],
            item["主体差异"] if item["主体差异"] == "—" else float(Decimal(item["主体差异"])),
            float(Decimal(item["对集团差异的影响"])),
        ])
        row_index = subject.max_row
        if item["主体差异"] != "—":
            subject.cell(row_index, 3).number_format = signed_number_format
        subject.cell(row_index, 4).number_format = signed_number_format

    instructions = workbook.create_sheet("恢复输入说明")
    instructions.append(["字段", "说明"])
    instructions.append(["reviewed_by", "审核人，必填"])
    instructions.append(["reason", "差异原因和修正依据，必填"])
    instructions.append(["corrected_revenue.current_month", "本年本月正确营业收入，仅有差异时填写"])
    instructions.append(["corrected_revenue.current_cumulative", "本年累计正确营业收入，仅有差异时填写"])
    instructions.append(["corrected_revenue.prior_month", "去年本月正确营业收入，仅有差异时填写"])
    instructions.append(["corrected_revenue.prior_cumulative", "去年累计正确营业收入，仅有差异时填写"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def prepare_difference_review(state: ProfitWorkflowState) -> dict[str, Any]:
    current_dataset, prior_dataset, current_ledger, _, *_ = _m7_inputs(state)
    rows, evidence = core.enrich_m7_rows_with_unbilled(
        [dict(row) for row in state["m7_differences"]],
        current_dataset, prior_dataset, state["year"], state["month"],
    )
    subject_rows = core.build_current_subject_m7_breakdown(
        _artifact_path(state, "current_m3"),
        _artifact_path(state, "current_m4_summary"),
        current_ledger,
        rows,
    )
    review_round = state.get("review_round", 0) + 1
    review_dir = _dirs(state)["m7"] / f"review_round_{review_round:02d}"
    subject_path = core.write_subject_m7_breakdown(review_dir, subject_rows)
    review_path = review_dir / "03_差异审核清单.xlsx"
    write_difference_review_workbook(review_path, rows, evidence, subject_rows, state, review_round)
    evidence_rows = _json_evidence(evidence)
    updates = {
        "workflow_status": "WAITING_REVIEW",
        "review_round": review_round,
        "m7_differences": rows,
        "subject_differences": subject_rows,
        "unbilled_evidence": evidence_rows,
        "difference_review_path": str(review_path),
        "artifacts": {
            f"review_round_{review_round:02d}": str(review_path),
            f"subject_breakdown_round_{review_round:02d}": str(subject_path),
        },
    }
    _write_status(state, updates)
    return updates


_CORRECTION_KEYS = {
    ("本年", "本月"): "current_month",
    ("本年", "累计"): "current_cumulative",
    ("去年同期", "本月"): "prior_month",
    ("去年同期", "累计"): "prior_cumulative",
}


def _validate_correction(state: ProfitWorkflowState, response: Any) -> dict[str, Any]:
    if not isinstance(response, dict):
        raise ValueError("恢复输入必须是JSON对象")
    reviewed_by = core.compact_text(response.get("reviewed_by") or response.get("reviewer"))
    reason = core.compact_text(response.get("reason"))
    values = response.get("corrected_revenue")
    if not reviewed_by or not reason or not isinstance(values, dict):
        raise ValueError("恢复输入必须包含reviewed_by、reason和corrected_revenue")
    normalized: dict[str, str | None] = {}
    for key in _CORRECTION_KEYS.values():
        raw = values.get(key)
        if raw is None or raw == "":
            normalized[key] = None
            continue
        amount = core.to_number(raw)
        if amount is None:
            raise ValueError(f"corrected_revenue.{key}不是合法金额")
        normalized[key] = format(amount, ".2f")
    for row in state["m7_differences"]:
        if core.elimination_to_money(row["差异"]) == core.ZERO:
            continue
        key = _CORRECTION_KEYS[(row["年度类型"], row["范围"])]
        if normalized[key] is None:
            raise ValueError(f"{row['年度类型']}{row['范围']}存在差异，必须填写corrected_revenue.{key}")
    return {"reviewed_by": reviewed_by, "reason": reason, "corrected_revenue": normalized}


def _verify_fingerprints(state: ProfitWorkflowState) -> None:
    current = _fingerprint(_current_fingerprint_paths(state), state["current_period_policy"])
    prior = _fingerprint(_prior_fingerprint_paths(state), state["prior_period_policy"])
    if current != state["current_input_fingerprint"] or prior != state["prior_input_fingerprint"]:
        raise ValueError("暂停期间输入文件或规则依据发生变化，旧Checkpoint不可继续使用")


def _apply_correct_values(
    rows: list[dict[str, str]], monthly: str | None, cumulative: str | None, review_round: int
) -> list[dict[str, str]]:
    revenue = next(row for row in rows if row["项目"] == "营业收入")
    if monthly is not None:
        revenue["抵销后本月"] = monthly
    if cumulative is not None:
        revenue["抵销后累计"] = cumulative
    revenue["说明"] = f"{revenue.get('说明', '')}；第{review_round}轮人工确认正确营业收入".strip("；")
    core.recalculate_operating_cost_after_elimination(rows)
    return rows


def _load_or_save_base(path: Path, rows: list[dict[str, str]]) -> list[dict[str, str]]:
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return copy.deepcopy(rows)


def review_and_correct_revenue(state: ProfitWorkflowState) -> dict[str, Any]:
    from langgraph.types import interrupt

    review_file = _reference_path(state["difference_review_path"])
    response = interrupt({
        "workflow_status": "WAITING_REVIEW",
        "review_round": state["review_round"],
        "review_file": state["difference_review_path"],
        "m7_differences": state["m7_differences"],
        "required_fields": [
            _CORRECTION_KEYS[(row["年度类型"], row["范围"])]
            for row in state["m7_differences"]
            if core.elimination_to_money(row["差异"]) != core.ZERO
        ],
    })
    correction = _validate_correction(state, response)
    _verify_fingerprints(state)

    review_dir = review_file.parent
    current_path = _artifact_path(state, "current_m5")
    prior_path = _artifact_path(state, "prior_m5")
    current_base_path = review_dir / "04_本年修正前M5.json"
    prior_base_path = review_dir / "05_去年修正前M5.json"
    current_base = _load_or_save_base(
        current_base_path, core.read_result_rows(current_path)
    )
    prior_base = _load_or_save_base(
        prior_base_path, core.read_result_rows(prior_path)
    )
    values = correction["corrected_revenue"]
    current_rows = _apply_correct_values(
        current_base, values["current_month"], values["current_cumulative"], state["review_round"]
    )
    prior_rows = _apply_correct_values(
        prior_base, values["prior_month"], values["prior_cumulative"], state["review_round"]
    )
    core.write_result_rows(current_path, current_rows)
    core.write_result_rows(prior_path, prior_rows)

    record = {
        "round": state["review_round"],
        "reviewed_by": correction["reviewed_by"],
        "reason": correction["reason"],
        "reviewed_at": datetime.now().isoformat(timespec="seconds"),
        "before_differences": state["m7_differences"],
        "before_revenue": {
            "current": {key: format(value, ".2f") for key, value in core.read_final_revenue_values(current_base).items()},
            "prior": {key: format(value, ".2f") for key, value in core.read_final_revenue_values(prior_base).items()},
        },
        "corrected_revenue": values,
    }
    record_path = review_dir / "06_人工修正记录.json"
    record_path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    updates = {
        "workflow_status": "RUNNING",
        "corrected_revenue": correction,
        "review_history": [record],
        "artifacts": {
            f"correction_round_{state['review_round']:02d}": str(record_path),
            f"review_base_current_round_{state['review_round']:02d}": str(current_base_path),
            f"review_base_prior_round_{state['review_round']:02d}": str(prior_base_path),
            "current_m5": str(current_path),
            "prior_m5": str(prior_path),
        },
    }
    _write_status(state, updates)
    return updates


def publish_report(state: ProfitWorkflowState) -> dict[str, Any]:
    if state.get("has_difference") or core.nonzero_m7_keys(state["m7_differences"]):
        raise RuntimeError("M7四组差异未全部归零，禁止生成正式报表")
    current_rows = core.read_result_rows(_artifact_path(state, "current_m5"))
    prior_rows = core.read_result_rows(_artifact_path(state, "prior_m5"))
    current_yangyuan, prior_yangyuan, current_us, prior_us = _self_sales_and_us(state)
    output = _dirs(state)["m6"] / f"02_长安仁恒{state['year']}年{state['month']}月合并报表（正式）.xls"
    core.write_final_template(
        Path(str(state["selected_files"]["template_file"])), output,
        current_rows, prior_rows, current_yangyuan, prior_yangyuan,
        state["year"], state["month"], current_us, prior_us,
    )
    updates = {
        "workflow_status": "SUCCEEDED",
        "formal_report": str(output),
        "error": None,
        "artifacts": {"formal_report": str(output)},
    }
    _write_status(state, updates)
    return updates


def build_current_year_subgraph(artifact_persister=None, artifact_resolver=None):
    from langgraph.graph import END, START, StateGraph

    builder = StateGraph(ProfitWorkflowState)
    builder.add_node("current_m2_parse", _wrap_node(current_m2_parse, artifact_persister, artifact_resolver))
    builder.add_node("current_invoice_revenue", _wrap_node(current_invoice_revenue, artifact_persister, artifact_resolver))
    builder.add_node("current_m3_pre_elimination", _wrap_node(current_m3_pre_elimination, artifact_persister, artifact_resolver))
    builder.add_node("current_m4_internal_sales", _wrap_node(current_m4_internal_sales, artifact_persister, artifact_resolver))
    builder.add_node("current_m5_final", _wrap_node(current_m5_final, artifact_persister, artifact_resolver))
    builder.add_edge(START, "current_m2_parse")
    builder.add_edge("current_m2_parse", "current_invoice_revenue")
    builder.add_edge("current_invoice_revenue", "current_m3_pre_elimination")
    builder.add_edge("current_m3_pre_elimination", "current_m4_internal_sales")
    builder.add_edge("current_m4_internal_sales", "current_m5_final")
    builder.add_edge("current_m5_final", END)
    return builder.compile()


def build_prior_year_subgraph(artifact_persister=None, artifact_resolver=None):
    from langgraph.graph import END, START, StateGraph

    builder = StateGraph(ProfitWorkflowState)
    builder.add_node("prior_m2_parse", _wrap_node(prior_m2_parse, artifact_persister, artifact_resolver))
    builder.add_node("prior_m3_pre_summary", _wrap_node(prior_m3_pre_summary, artifact_persister, artifact_resolver))
    builder.add_node("prior_m4_ledger_components", _wrap_node(prior_m4_ledger_components, artifact_persister, artifact_resolver))
    builder.add_node("prior_m5_final", _wrap_node(prior_m5_final, artifact_persister, artifact_resolver))
    builder.add_edge(START, "prior_m2_parse")
    builder.add_edge("prior_m2_parse", "prior_m3_pre_summary")
    builder.add_edge("prior_m3_pre_summary", "prior_m4_ledger_components")
    builder.add_edge("prior_m4_ledger_components", "prior_m5_final")
    builder.add_edge("prior_m5_final", END)
    return builder.compile()


def build_workflow(checkpointer=None, artifact_persister=None, artifact_resolver=None):
    from langgraph.graph import END, START, StateGraph

    builder = StateGraph(ProfitWorkflowState)
    builder.add_node("determine_period_policy", _wrap_node(determine_period_policy, artifact_persister, artifact_resolver))
    builder.add_node("preflight", _wrap_node(preflight, artifact_persister, artifact_resolver))
    builder.add_node("prepare_run", _wrap_node(prepare_run, artifact_persister, artifact_resolver))
    builder.add_node(
        "current_year_subgraph",
        build_current_year_subgraph(artifact_persister, artifact_resolver),
    )
    builder.add_node(
        "prior_year_subgraph",
        build_prior_year_subgraph(artifact_persister, artifact_resolver),
    )
    builder.add_node("m6_preview", _wrap_node(m6_preview, artifact_persister, artifact_resolver))
    builder.add_node("m7_group_validation", _wrap_node(m7_group_validation, artifact_persister, artifact_resolver))
    builder.add_node("prepare_difference_review", _wrap_node(prepare_difference_review, artifact_persister, artifact_resolver))
    builder.add_node("review_and_correct_revenue", _wrap_node(review_and_correct_revenue, artifact_persister, artifact_resolver))
    builder.add_node("publish_report", _wrap_node(publish_report, artifact_persister, artifact_resolver))

    builder.add_edge(START, "determine_period_policy")
    builder.add_edge("determine_period_policy", "preflight")
    builder.add_edge("preflight", "prepare_run")
    builder.add_edge("prepare_run", "current_year_subgraph")
    builder.add_edge("current_year_subgraph", "prior_year_subgraph")
    builder.add_edge("prior_year_subgraph", "m6_preview")
    builder.add_edge("m6_preview", "m7_group_validation")
    builder.add_conditional_edges(
        "m7_group_validation", route_after_m7,
        {"prepare_difference_review": "prepare_difference_review", "publish_report": "publish_report"},
    )
    builder.add_edge("prepare_difference_review", "review_and_correct_revenue")
    builder.add_edge("review_and_correct_revenue", "m6_preview")
    builder.add_edge("publish_report", END)
    return builder.compile(checkpointer=checkpointer)
